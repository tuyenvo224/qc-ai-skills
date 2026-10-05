import argparse
import re
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

INVALID_SHEET_CHARS = re.compile(r'[\[\]\*\?/\\:]')


def clean_cell(text: str) -> str:
    text = text.strip()
    text = re.sub(r'<br\s*/?>', '\n', text, flags=re.IGNORECASE)
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'(?<!\w)\*(.+?)\*(?!\w)', r'\1', text)
    text = re.sub(r'`([^`]+)`', r'\1', text)
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    text = text.replace('\\|', '|')
    return text.strip()


def split_row(line: str):
    line = line.strip()
    if line.startswith('|'):
        line = line[1:]
    if line.endswith('|'):
        line = line[:-1]
    cells = re.split(r'(?<!\\)\|', line)
    return [clean_cell(c) for c in cells]


def is_separator(line: str) -> bool:
    stripped = line.strip()
    if not stripped or '-' not in stripped:
        return False
    return bool(re.match(r'^[|\s\-:]+$', stripped))


def find_tables(md_text: str):
    lines = md_text.splitlines()
    tables = []
    current_heading = None
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        heading_match = re.match(r'^(#{1,6})\s+(.*)', line)
        if heading_match:
            current_heading = heading_match.group(2).strip()
            i += 1
            continue
        if line.strip().startswith('|') and i + 1 < n and is_separator(lines[i + 1]):
            header = split_row(line)
            j = i + 2
            rows = []
            while j < n and lines[j].strip().startswith('|'):
                rows.append(split_row(lines[j]))
                j += 1
            tables.append({'heading': current_heading, 'header': header, 'rows': rows})
            i = j
            continue
        i += 1
    return tables


def safe_sheet_name(name: str, used: set) -> str:
    name = INVALID_SHEET_CHARS.sub('', name)
    name = re.sub(r'\s+', ' ', name).strip() or 'Sheet'
    name = name[:31]
    base, idx = name, 2
    while name in used:
        suffix = f'_{idx}'
        name = base[:31 - len(suffix)] + suffix
        idx += 1
    used.add(name)
    return name


def autofit_columns(ws, max_width=60):
    widths = {}
    for row in ws.iter_rows():
        for cell in row:
            if cell.value is None:
                continue
            longest = max(len(l) for l in str(cell.value).split('\n'))
            widths[cell.column_letter] = max(widths.get(cell.column_letter, 0), longest)
    for col, width in widths.items():
        ws.column_dimensions[col].width = min(max(width + 2, 10), max_width)


def write_table_sheet(wb, name, header, rows, used_names):
    ws = wb.create_sheet(safe_sheet_name(name, used_names))
    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(start_color='4472C4', end_color='4472C4', fill_type='solid')
    header_align = Alignment(wrap_text=True, vertical='center', horizontal='center')
    body_align = Alignment(wrap_text=True, vertical='top')
    thin = Side(style='thin', color='D9D9D9')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)

    ws.append(header)
    for cell in ws[1]:
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = border

    col_count = len(header)
    for row in rows:
        if len(row) < col_count:
            row = row + [''] * (col_count - len(row))
        elif len(row) > col_count:
            row = row[:col_count]
        ws.append(row)

    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = body_align
            cell.border = border

    ws.freeze_panes = 'A2'
    ws.auto_filter.ref = ws.dimensions
    autofit_columns(ws)
    return ws


def write_plain_text_sheet(wb, name, md_text, used_names):
    ws = wb.create_sheet(safe_sheet_name(name, used_names))
    ws.append(['Content'])
    ws['A1'].font = Font(bold=True)
    for line in md_text.splitlines():
        cleaned = clean_cell(line)
        if cleaned:
            ws.append([cleaned])
    autofit_columns(ws)
    return ws


def main():
    parser = argparse.ArgumentParser(description='Convert Markdown tables to an Excel workbook (.xlsx)')
    parser.add_argument('input', help='Path to the source .md file')
    parser.add_argument('output', nargs='?', help='Path to the output .xlsx file (default: alongside input, same name)')
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        print(f'Input file not found: {input_path}', file=sys.stderr)
        sys.exit(1)

    output_path = Path(args.output) if args.output else input_path.with_suffix('.xlsx')

    md_text = input_path.read_text(encoding='utf-8')
    tables = find_tables(md_text)

    wb = Workbook()
    wb.remove(wb.active)
    used_names = set()

    if tables:
        for idx, table in enumerate(tables, start=1):
            default_name = input_path.stem if len(tables) == 1 else f'Table{idx}'
            name = table['heading'] or default_name
            write_table_sheet(wb, name, table['header'], table['rows'], used_names)
    else:
        write_plain_text_sheet(wb, input_path.stem, md_text, used_names)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f'Da tao file Excel: {output_path}')


if __name__ == '__main__':
    main()
