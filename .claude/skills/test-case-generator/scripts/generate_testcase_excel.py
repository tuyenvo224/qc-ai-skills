"""
Generate the formatted Test Case Excel workbook required by the test-case-generator skill
(see ../references/excel-export.md for the full formatting spec this script implements).

Usage:
    python generate_testcase_excel.py <input.json> <output.xlsx>

Input JSON schema:
{
  "project_name": "GiftPort",
  "feature_name": "Category Management",
  "version": "v1.0",                      # optional, default "v1.0"
  "test_cases": [
    {
      "id": "TC-CAT-001",
      "technique": "UC",                  # one of EP,BVA,DT,ST,UC,PW,EG,CL,EXP
      "priority": "P2",                   # P1..P4
      "objective": "[UC] Verify danh sach displays correctly",
      "steps": "1. Navigate to list screen\\n2. Observe grid",
      "expected": "Grid shows 10 records\\nSorted DESC by created_at",
      "requirement": "REQ-001, REQ-002",  # optional — comma-separated REQ-xxx ids (see compute_coverage.py)
      "risk": "",                         # optional
      "status": "",                       # optional, Pass/Fail/Blocked/'' (Not Run)
      "test_type": "FN,UI"                # optional, comma-separated: FN,BL,NEG,EC,PM,DI,ST,CN,IT,SEC,UI
    }
  ],
  "metadata": {                           # optional — the 8-row company-form header block (rows 1-8)
                                           # is ALWAYS written above the title/summary/header rows on
                                           # the "Test Cases" sheet (title/summary/header/data shift
                                           # down by 8 rows accordingly, regardless of this key).
                                           # This object only lets you fill in the block's values;
                                           # omit it (or omit any field) and that cell is left blank
                                           # ("screen_name" falls back to top-level "feature_name").
      "screen_name": "Tên màn hình/chức năng",
      "spec_link": "",
      "redmine_link": "",
      "prototype": "",
      "domain": "",
      "account": "",
      "qc_name": "",                       # "QC thực hiện" - người tạo TCs
      "qc_date": "",                       # optional, default = today (YYYY-MM-DD)
      "work_type": "Tạo TCs"               # optional, default "Tạo TCs"
    },
  "assumptions": [                        # optional — Sheet 3 only created if non-empty
    {
      "id": "ASM-001",
      "category": "Business Rule",
      "description": "...",
      "impact": "Medium",
      "test_cases_affected": "TC-001, TC-022",
      "status": "Pending",                # Pending/Confirmed/Rejected/Modified
      "verified_by": "",
      "date": "",
      "resolution": ""
    }
  ]
}
"""

import argparse
import json
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

TECHNIQUES = ["EP", "BVA", "DT", "ST", "UC", "PW", "EG", "CL", "EXP"]
TECHNIQUE_LABELS = {
    "EP": "[EP] Equivalence Partitioning",
    "BVA": "[BVA] Boundary Value Analysis",
    "DT": "[DT] Decision Table",
    "ST": "[ST] State Transition",
    "UC": "[UC] Use Case",
    "PW": "[PW] Pairwise",
    "EG": "[EG] Error Guessing",
    "CL": "[CL] Checklist-Based",
    "EXP": "[EXP] Exploratory",
}
TEST_TYPES = ["FN", "BL", "NEG", "EC", "PM", "DI", "ST", "CN", "IT", "SEC", "UI", "A11Y", "COMPAT"]
TEST_TYPE_LABELS = {
    "FN": "Functional (FN)",
    "BL": "Business Logic (BL)",
    "NEG": "Negative (NEG)",
    "EC": "Edge Cases (EC)",
    "PM": "Permission (PM)",
    "DI": "Data Integrity (DI)",
    "ST": "State Testing (ST)",
    "CN": "Concurrent (CN)",
    "IT": "Integration (IT)",
    "SEC": "Security (SEC)",
    "UI": "UI Basic (UI)",
    "A11Y": "Accessibility (A11Y)",
    "COMPAT": "Compatibility (COMPAT)",
}
PRIORITIES = ["P1", "P2", "P3", "P4"]

BORDER = Border(
    left=Side(style="thin"), right=Side(style="thin"),
    top=Side(style="thin"), bottom=Side(style="thin"),
)

PRIORITY_STYLE = {
    "P1": dict(bg="FF0000", font=Font(bold=True, color="FFFFFF", size=10)),
    "P2": dict(bg="FFC000", font=Font(bold=True, color="000000", size=10)),
    "P3": dict(bg="92D050", font=Font(bold=True, color="000000", size=10)),
    "P4": dict(bg="D9D9D9", font=Font(color="000000", size=10)),
}
STATUS_STYLE = {
    "Pass": dict(bg="C6EFCE", font=Font(bold=True, color="006100", size=10)),
    "Fail": dict(bg="FFC7CE", font=Font(bold=True, color="9C0006", size=10)),
    "Blocked": dict(bg="FFEB9C", font=Font(bold=True, color="9C5700", size=10)),
    "Deprecated": dict(bg="D9D9D9", font=Font(italic=True, color="595959", size=10)),
    "": dict(bg="F2F2F2", font=Font(color="595959", size=10)),
}

HEADERS = ["STT", "Test ID", "Technique", "Priority", "Test Objective",
           "Test Steps", "Expected Result", "Requirement", "Risk/Bug", "Status"]
COL_WIDTHS = [5, 15, 10, 8, 38, 48, 48, 14, 15, 12]


def fill(hex_color):
    return PatternFill(start_color=hex_color, end_color=hex_color, fill_type="solid")


def parse_test_types(tc):
    raw = tc.get("test_type", "") or ""
    return [t.strip().upper() for t in raw.split(",") if t.strip()]


def _wrapped_line_count(text, col_width_chars):
    """Estimate how many display lines `text` occupies once wrapped at
    `col_width_chars` characters — openpyxl/Excel cannot autofit row height for
    wrapped text on save, so row height must be computed manually."""
    if not text:
        return 1
    width = max(1, int(col_width_chars))
    total = 0
    for line in str(text).split("\n"):
        total += max(1, -(-len(line) // width))  # ceil division
    return total


def _autofit_row_height(tc):
    obj_lines = _wrapped_line_count(tc.get("objective", ""), COL_WIDTHS[4])
    steps_lines = _wrapped_line_count(tc.get("steps", ""), COL_WIDTHS[5])
    expected_lines = _wrapped_line_count(tc.get("expected", ""), COL_WIDTHS[6])
    risk_lines = _wrapped_line_count(tc.get("risk", ""), COL_WIDTHS[8])
    max_lines = max(obj_lines, steps_lines, expected_lines, risk_lines)
    return max(35, min(500, max_lines * 15 + 10))


def _parse_qc_date(value):
    if not value:
        return datetime.now()
    if isinstance(value, datetime):
        return value
    for fmt in ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(str(value), fmt)
        except ValueError:
            continue
    return datetime.now()


def _write_metadata_header(ws, metadata):
    """Write the 8-row company-form header block (rows 1-8) matching
    assets/Template_TCs.xlsx: a title row, a 'screen name / spec link /
    redmine / prototype / domain / account' form in columns E:F, and a
    'QC thực hiện' sub-block (who/when/what) in columns H:J."""
    label_font = Font(bold=True, size=11)

    ws["E1"] = "KỊCH BẢN KIỂM THỬ *"
    ws["E1"].font = Font(bold=True, size=20)
    ws["E1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 26

    def label(coord, text):
        cell = ws[coord]
        cell.value = text
        cell.font = label_font
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center")

    def value(coord, val, number_format=None):
        cell = ws[coord]
        cell.value = val
        cell.border = BORDER
        cell.alignment = Alignment(vertical="center")
        if number_format:
            cell.number_format = number_format

    label("E2", "Tên màn hình/Tên chức năng")
    value("F2", metadata.get("screen_name", ""))

    yellow = fill("FFFF00")
    for coord in ("H2", "I2", "J2"):
        ws[coord].fill = yellow
        ws[coord].border = BORDER
    ws["H2"].value = "QC thực hiện"
    ws["H2"].font = label_font

    label("E3", "Link spec")
    value("F3", metadata.get("spec_link", ""))
    value("H3", metadata.get("qc_name", ""))
    value("I3", _parse_qc_date(metadata.get("qc_date")), number_format="mm-dd-yy")
    value("J3", metadata.get("work_type") or "Tạo TCs")
    ws.row_dimensions[3].height = 30

    label("E4", "Link redmine")
    value("F4", metadata.get("redmine_link", ""))

    label("E5", "Prototype")
    value("F5", metadata.get("prototype", ""))

    label("E6", "Domain")
    value("F6", metadata.get("domain", ""))

    label("E7", "Account")
    value("F7", metadata.get("account", ""))

    for coord in ("E8", "F8", "H8", "I8", "J8"):
        ws[coord].border = BORDER
    ws.row_dimensions[8].height = 15


def build_test_cases_sheet(wb, data):
    test_cases = data["test_cases"]
    project_name = data.get("project_name", "")
    feature_name = data.get("feature_name", "")
    metadata = dict(data.get("metadata") or {})
    metadata.setdefault("screen_name", feature_name)
    offset = 8

    ws = wb.active
    ws.title = "Test Cases"

    counts = {p: sum(1 for tc in test_cases if tc.get("priority") == p) for p in PRIORITIES}

    _write_metadata_header(ws, metadata)

    title_row = 1 + offset
    meta_row = 2 + offset
    header_row = 3 + offset

    ws.merge_cells(f"A{title_row}:J{title_row}")
    ws[f"A{title_row}"] = f"{project_name} - {feature_name} - Test Cases"
    ws[f"A{title_row}"].font = Font(bold=True, size=14, color="FFFFFF")
    ws[f"A{title_row}"].fill = fill("203864")
    ws[f"A{title_row}"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[title_row].height = 25

    ws.merge_cells(f"A{meta_row}:J{meta_row}")
    ws[f"A{meta_row}"] = (
        f"Generated: {datetime.now().strftime('%Y-%m-%d')} | "
        f"Total: {len(test_cases)} | P1: {counts['P1']} | P2: {counts['P2']} | "
        f"P3: {counts['P3']} | P4: {counts['P4']}"
    )
    ws[f"A{meta_row}"].font = Font(italic=True, size=10)
    ws[f"A{meta_row}"].alignment = Alignment(horizontal="center")
    ws.row_dimensions[meta_row].height = 18

    header_fill = fill("4472C4")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    for col_num, header in enumerate(HEADERS, 1):
        cell = ws.cell(row=header_row, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[header_row].height = 30

    for col_num, width in enumerate(COL_WIDTHS, 1):
        ws.column_dimensions[get_column_letter(col_num)].width = width

    for idx, tc in enumerate(test_cases, start=1):
        row_num = idx + header_row
        priority = tc.get("priority", "")
        status_val = (tc.get("status") or "").strip()
        technique = tc.get("technique", "")

        ws.append([
            idx,
            tc.get("id", ""),
            f"[{technique}]" if technique else "",
            priority,
            tc.get("objective", ""),
            tc.get("steps", ""),
            tc.get("expected", ""),
            tc.get("requirement", ""),
            tc.get("risk", ""),
            status_val,
        ])

        for col_num in range(1, 11):
            cell = ws.cell(row=row_num, column=col_num)
            cell.border = BORDER
            horizontal = "center" if col_num in (1, 2, 3, 4, 10) else "left"
            cell.alignment = Alignment(horizontal=horizontal, vertical="top", wrap_text=True)

            if col_num == 4 and priority in PRIORITY_STYLE:
                cell.fill = fill(PRIORITY_STYLE[priority]["bg"])
                cell.font = PRIORITY_STYLE[priority]["font"]

            if col_num == 10:
                style = STATUS_STYLE.get(status_val, STATUS_STYLE[""])
                cell.fill = fill(style["bg"])
                cell.font = style["font"]

        ws.row_dimensions[row_num].height = _autofit_row_height(tc)

    last_data_row = len(test_cases) + header_row
    first_data_row = header_row + 1
    status_range = f"J{first_data_row}:J{last_data_row}"
    ws.conditional_formatting.add(status_range, FormulaRule(formula=[f'$J{first_data_row}="Pass"'], fill=fill("C6EFCE")))
    ws.conditional_formatting.add(status_range, FormulaRule(formula=[f'$J{first_data_row}="Fail"'], fill=fill("FFC7CE")))
    ws.conditional_formatting.add(status_range, FormulaRule(formula=[f'$J{first_data_row}="Blocked"'], fill=fill("FFEB9C")))
    ws.conditional_formatting.add(status_range, FormulaRule(formula=[f'$J{first_data_row}="Deprecated"'], fill=fill("D9D9D9")))
    ws.conditional_formatting.add(status_range, FormulaRule(formula=[f'$J{first_data_row}=""'], fill=fill("F2F2F2")))

    ws.freeze_panes = f"E{first_data_row}"
    ws.auto_filter.ref = f"A{header_row}:J{last_data_row}"
    return counts


def build_summary_sheet(wb, data, priority_counts):
    test_cases = data["test_cases"]
    ws = wb.create_sheet("Summary")
    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 14

    bold = Font(bold=True)
    block_header_fill = fill("D9E1F2")

    ws.merge_cells("A1:B1")
    ws["A1"] = "Project Summary"
    ws["A1"].font = Font(bold=True, size=14)

    ws["A3"], ws["B3"] = "Project Name", data.get("project_name", "")
    ws["A4"], ws["B4"] = "Feature/Module", data.get("feature_name", "")
    ws["A5"], ws["B5"] = "Generated Date", datetime.now().strftime("%Y-%m-%d")
    ws["A6"], ws["B6"] = "Version", data.get("version", "v1.0")
    ws["A7"], ws["B7"] = "Total Test Cases", len(test_cases)
    for r in range(3, 8):
        ws.cell(row=r, column=1).font = bold

    row = 9
    ws.cell(row=row, column=1, value="By Priority").font = bold
    ws.cell(row=row, column=2, value="Count").font = bold
    ws.cell(row=row, column=1).fill = block_header_fill
    ws.cell(row=row, column=2).fill = block_header_fill
    row += 1
    priority_labels = {"P1": "P1 (Critical)", "P2": "P2 (High)", "P3": "P3 (Medium)", "P4": "P4 (Low)"}
    for p in PRIORITIES:
        ws.cell(row=row, column=1, value=priority_labels[p])
        ws.cell(row=row, column=2, value=priority_counts.get(p, 0))
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="By Technique").font = bold
    ws.cell(row=row, column=2, value="Count").font = bold
    ws.cell(row=row, column=1).fill = block_header_fill
    ws.cell(row=row, column=2).fill = block_header_fill
    row += 1
    for tag in TECHNIQUES:
        count = sum(1 for tc in test_cases if (tc.get("technique") or "").upper() == tag)
        ws.cell(row=row, column=1, value=TECHNIQUE_LABELS[tag])
        ws.cell(row=row, column=2, value=count)
        row += 1

    row += 1
    ws.cell(row=row, column=1, value="By Test Type").font = bold
    ws.cell(row=row, column=2, value="Count").font = bold
    ws.cell(row=row, column=1).fill = block_header_fill
    ws.cell(row=row, column=2).fill = block_header_fill
    row += 1
    any_test_type = any(parse_test_types(tc) for tc in test_cases)
    if any_test_type:
        for t in TEST_TYPES:
            count = sum(1 for tc in test_cases if t in parse_test_types(tc))
            ws.cell(row=row, column=1, value=TEST_TYPE_LABELS[t])
            ws.cell(row=row, column=2, value=count)
            row += 1
    else:
        ws.cell(row=row, column=1, value="(test_type not provided in input data)")
        row += 1


def build_assumptions_sheet(wb, assumptions):
    if not assumptions:
        return
    ws = wb.create_sheet("Assumptions")
    headers = ["Assumption ID", "Category", "Description", "Impact",
               "Test Cases Affected", "Status", "Verified By", "Date", "Resolution"]
    widths = [12, 15, 40, 10, 20, 12, 15, 12, 30]
    header_fill = fill("4472C4")
    header_font = Font(bold=True, color="FFFFFF")

    for col_num, (header, width) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col_num, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_num)].width = width

    status_fill = {"Pending": fill("FFFF00"), "Confirmed": fill("00FF00"), "Rejected": fill("FF0000")}

    keys = ["id", "category", "description", "impact", "test_cases_affected",
            "status", "verified_by", "date", "resolution"]
    for idx, item in enumerate(assumptions, start=2):
        for col_num, key in enumerate(keys, 1):
            cell = ws.cell(row=idx, column=col_num, value=item.get(key, ""))
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if key == "status" and item.get("status") in status_fill:
                cell.fill = status_fill[item["status"]]

        desc_lines = _wrapped_line_count(item.get("description", ""), widths[2])
        resolution_lines = _wrapped_line_count(item.get("resolution", ""), widths[8])
        ws.row_dimensions[idx].height = max(20, min(300, max(desc_lines, resolution_lines) * 15 + 10))

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:I{len(assumptions) + 1}"


def create_testcase_excel(data, output_file):
    wb = Workbook()
    priority_counts = build_test_cases_sheet(wb, data)
    build_summary_sheet(wb, data, priority_counts)
    build_assumptions_sheet(wb, data.get("assumptions", []))
    wb.save(output_file)


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("input_json", help="Path to input JSON file (see schema in this script's docstring)")
    parser.add_argument("output_xlsx", nargs="?", help="Path to output .xlsx file (default: same name as input, .xlsx)")
    args = parser.parse_args()

    input_path = Path(args.input_json)
    if not input_path.exists():
        print(f"Error: input file not found: {input_path}", file=sys.stderr)
        sys.exit(1)

    output_path = Path(args.output_xlsx) if args.output_xlsx else input_path.with_suffix(".xlsx")

    data = json.loads(input_path.read_text(encoding="utf-8"))
    if not data.get("test_cases"):
        print("Error: input JSON must contain a non-empty 'test_cases' array", file=sys.stderr)
        sys.exit(1)

    create_testcase_excel(data, str(output_path))
    print(f"Created {output_path} ({len(data['test_cases'])} test cases)")


if __name__ == "__main__":
    main()
