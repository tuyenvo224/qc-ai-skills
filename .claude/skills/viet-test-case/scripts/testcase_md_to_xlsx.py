"""
Convert the Markdown test case table produced by the viet-test-case skill into the
company's standard "KICH BAN KIEM THU" Excel file, using
../assets/Template_TCs.xlsx as the base (sheet "TCs") so the output keeps the exact
corporate layout: a metadata header block (feature name, spec/redmine links, QC info)
followed by test cases grouped into "I. <module>", "II. <module>"... blocks.

Expected Markdown table columns (order-independent, matched by keyword):
    STT | Module/Function | Muc dich kiem thu | Precondition | Test Data | Cac buoc thuc hien | Ket qua mong muon

- "Module/Function" groups consecutive rows into a Group block in the output.
- "Type" (Positive/Negative/Boundary/Security) is expected as a "[Type] " prefix
  inside "Muc dich kiem thu" - the template has no dedicated Type column.

The 4 execution-only columns in the template (Ket qua hien tai, Ma loi, Ghi chu,
QC thuc hien) are left blank - they are filled in by the QC during test execution,
not authored here.

Table-parsing logic (clean_cell/split_row/is_separator/find_tables) is adapted from
../../markdown-to-excel/scripts/md_to_xlsx.py - keep both in sync if the generic
parsing rules change.

Usage:
    python testcase_md_to_xlsx.py <input.md> [output.xlsx] [--qc-name "Ten QC"]
"""

import argparse
import re
import sys
import unicodedata
from copy import copy
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill

TEMPLATE_PATH = Path(__file__).resolve().parent.parent / 'assets' / 'Template_TCs.xlsx'
SHEET_NAME = 'TCs'

# Rows in the template's example content (3 sample groups, "I."/"II."/"III.") that
# get wiped and rebuilt from the parsed Markdown data.
EXAMPLE_FIRST_ROW = 12
EXAMPLE_ROW_COUNT = 16  # rows 12..27 inclusive
EXAMPLE_MERGES = ['A12:E12', 'A19:E19', 'A27:E27']
GROUP_MERGE_COLS = 5  # group header spans columns A..E

DATA_COLS = 10  # A..J
EXPECTED_COL = 6  # F - "Ket qua mong muon"

# Expected results that rest on an assumption / still need PO-BA confirmation get a
# light-yellow fill so the QC can spot them at a glance. Matched against the
# normalize()d text (no diacritics, lowercase).
CONFIRM_FILL = PatternFill(fill_type='solid', start_color='FFFFFF99', end_color='FFFFFF99')
CONFIRM_PATTERN = re.compile(
    r'\b('
    r'can confirm|confirm lai|cho confirm|can xac nhan|cho xac nhan'
    r'|assumption|gia dinh|gia su|tam dung|tam hieu'
    r'|chua chot|chua ro|tbd'
    r'|(asm|qh|qm|ql)-\d+'
    r')\b'
)


def clean_cell(text: str) -> str:
    text = text.strip()
    text = re.sub(r'<br\s*/?>', '\n', text, flags=re.IGNORECASE)
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'(?<!\w)\*(.+?)\*(?!\w)', r'\1', text)
    text = re.sub(r'`([^`]+)`', r'\1', text)
    text = re.sub(r'\[([^\]]+)\]\(([^)]+)\)', r'\1 (\2)', text)
    text = text.replace('\\|', '|')
    return text.strip()


def split_row(line: str):
    line = line.strip()
    if line.startswith('|'):
        line = line[1:]
    if line.endswith('|'):
        line = line[:-1]
    cells = re.split(r'(?<!\\)\|', line)
    return [clean_cell(c) for c in cells]


def is_separator(line: str) -> bool:
    stripped = line.strip()
    if not stripped or '-' not in stripped:
        return False
    return bool(re.match(r'^[|\s\-:]+$', stripped))


def find_tables(md_text: str):
    lines = md_text.splitlines()
    tables = []
    current_heading = None
    i, n = 0, len(lines)
    while i < n:
        line = lines[i]
        heading_match = re.match(r'^(#{1,6})\s+(.*)', line)
        if heading_match:
            current_heading = heading_match.group(2).strip()
            i += 1
            continue
        if line.strip().startswith('|') and i + 1 < n and is_separator(lines[i + 1]):
            header = split_row(line)
            j = i + 2
            rows = []
            while j < n and lines[j].strip().startswith('|'):
                rows.append(split_row(lines[j]))
                j += 1
            tables.append({'heading': current_heading, 'header': header, 'rows': rows})
            i = j
            continue
        i += 1
    return tables


def normalize(text: str) -> str:
    text = text.replace('đ', 'd').replace('Đ', 'D')
    text = unicodedata.normalize('NFKD', text)
    text = ''.join(c for c in text if not unicodedata.combining(c))
    return text.lower().strip()


def needs_confirmation(text) -> bool:
    """True when an expected result is based on an assumption or is marked as
    needing PO/BA confirmation (e.g. "(Theo ASM-004 - can confirm QH-003)")."""
    return bool(text) and bool(CONFIRM_PATTERN.search(normalize(str(text))))


def locate_columns(header):
    normalized = [normalize(h) for h in header]

    def find(*keywords):
        for i, h in enumerate(normalized):
            if any(kw in h for kw in keywords):
                return i
        return None

    idx = {
        'stt': find('stt'),
        'module': find('module', 'function', 'chuc nang con'),
        'purpose': find('muc dich'),
        'precondition': find('precondition', 'dieu kien tien quyet'),
        'test_data': find('test data', 'du lieu'),
        'steps': find('buoc thuc hien', 'cac buoc'),
        'expected': find('ket qua mong muon', 'expected'),
    }
    missing = [k for k, v in idx.items() if v is None]
    if missing:
        raise ValueError(
            f"Khong tim thay cot: {missing}. Header nhan duoc: {header}"
        )
    return idx


def group_rows(rows, module_idx):
    groups = []
    for row in rows:
        module_name = row[module_idx] if module_idx < len(row) else ''
        if groups and groups[-1][0] == module_name:
            groups[-1][1].append(row)
        else:
            groups.append([module_name, [row]])
    return groups


def to_roman(n: int) -> str:
    values = [
        (1000, 'M'), (900, 'CM'), (500, 'D'), (400, 'CD'),
        (100, 'C'), (90, 'XC'), (50, 'L'), (40, 'XL'),
        (10, 'X'), (9, 'IX'), (5, 'V'), (4, 'IV'), (1, 'I'),
    ]
    result = []
    for value, symbol in values:
        count, n = divmod(n, value)
        result.append(symbol * count)
    return ''.join(result)


def extract_feature_info(md_text: str, input_path: Path):
    heading = ''
    for line in md_text.splitlines():
        m = re.match(r'^#\s+(.*)', line)
        if m:
            heading = m.group(1).strip()
            break
    if not heading:
        heading = input_path.stem

    heading = re.sub(r'^Test Case:\s*', '', heading, flags=re.IGNORECASE).strip()

    redmine_url = None
    url_match = re.search(r'https?://\S*redmine\S*', heading)
    if url_match:
        redmine_url = url_match.group(0)
    else:
        num_match = re.search(r'#(\d{3,})', heading)
        if num_match:
            redmine_url = f'https://redmine.gotit.vn/issues/{num_match.group(1)}'

    return heading, redmine_url


def capture_row_style(ws, row):
    style = {}
    for col in range(1, DATA_COLS + 1):
        cell = ws.cell(row=row, column=col)
        style[col] = dict(
            font=copy(cell.font),
            fill=copy(cell.fill),
            border=copy(cell.border),
            alignment=copy(cell.alignment),
        )
    return style


def plain_text_font(font):
    """The template's sample cell in column F carries the built-in "Hyperlink"
    style (Calibri, blue theme color, underline), which makes test case text look
    like a clickable link. Force data cells to plain black, non-underlined text in
    the template's body font."""
    return Font(
        name='Times New Roman',
        sz=font.sz or 11,
        bold=font.b,
        italic=font.i,
        underline=None,
        strike=False,
        color='FF000000',
    )


def normalize_data_style(style):
    for s in style.values():
        s['font'] = plain_text_font(s['font'])
    return style


def apply_row_style(ws, row, style):
    for col in range(1, DATA_COLS + 1):
        cell = ws.cell(row=row, column=col)
        s = style[col]
        cell.font = s['font']
        cell.fill = s['fill']
        cell.border = s['border']
        cell.alignment = s['alignment']


MAX_ROW_HEIGHT = 409  # Excel's hard limit for a row height (points)


def wrapped_line_count(text, col_width):
    """Estimate how many visual lines `text` occupies in a column of `col_width`
    (Excel width units ~= characters), counting both explicit newlines and the
    soft wraps Excel applies to long lines when wrap_text is on."""
    # ~10% safety margin: Vietnamese diacritics/uppercase render wider than '0'.
    chars_per_line = max(1, int(col_width * 0.9))
    total = 0
    for line in str(text).split('\n'):
        total += max(1, -(-len(line) // chars_per_line))
    return total


def fit_row_to_text(ws, row):
    """Turn on wrap_text for every data cell in `row` and grow the row height so
    the full text of the tallest cell is visible (no clipped content)."""
    max_height = 15
    for col in range(1, DATA_COLS + 1):
        cell = ws.cell(row=row, column=col)
        alignment = copy(cell.alignment)
        alignment.wrap_text = True
        cell.alignment = alignment
        if cell.value in (None, ''):
            continue
        letter = cell.column_letter
        width = ws.column_dimensions[letter].width or 8.43
        font_size = cell.font.sz or 11
        line_height = font_size * 1.4
        lines = wrapped_line_count(cell.value, width)
        max_height = max(max_height, lines * line_height + 4)
    ws.row_dimensions[row].height = min(MAX_ROW_HEIGHT, max_height)


def build_workbook(header, rows, feature_name, redmine_url, qc_name):
    idx = locate_columns(header)
    groups = group_rows(rows, idx['module'])

    wb = openpyxl.load_workbook(TEMPLATE_PATH)
    ws = wb[SHEET_NAME]

    ws['F2'] = feature_name
    if redmine_url:
        ws['F4'] = redmine_url
    ws['H3'] = qc_name
    ws['I3'] = date.today()
    ws['J3'] = 'Tạo TCs'

    for col in ('C', 'D'):
        ws.column_dimensions[col].width = 35

    group_style = capture_row_style(ws, 12)
    data_style = normalize_data_style(capture_row_style(ws, 13))
    data_last_style = normalize_data_style(capture_row_style(ws, 18))

    for rng in EXAMPLE_MERGES:
        try:
            ws.unmerge_cells(rng)
        except KeyError:
            pass
    ws.delete_rows(EXAMPLE_FIRST_ROW, EXAMPLE_ROW_COUNT)

    total_rows_needed = sum(1 + len(items) for _, items in groups) if groups else 1
    ws.insert_rows(EXAMPLE_FIRST_ROW, amount=total_rows_needed)

    def cell_value(row, key):
        i = idx[key]
        return row[i] if i < len(row) else ''

    r = EXAMPLE_FIRST_ROW
    stt = 1
    highlighted = []
    for gi, (module_name, items) in enumerate(groups, start=1):
        apply_row_style(ws, r, group_style)
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=GROUP_MERGE_COLS)
        ws.cell(row=r, column=1, value=f'{to_roman(gi)}. {module_name or "Khac"}')
        r += 1
        for i, row in enumerate(items):
            is_last = i == len(items) - 1
            apply_row_style(ws, r, data_last_style if is_last else data_style)
            purpose = cell_value(row, 'purpose')
            precondition = cell_value(row, 'precondition')
            test_data = cell_value(row, 'test_data')
            steps = cell_value(row, 'steps')
            expected = cell_value(row, 'expected')
            ws.cell(row=r, column=1, value=stt)
            ws.cell(row=r, column=2, value=purpose)
            ws.cell(row=r, column=3, value=precondition)
            ws.cell(row=r, column=4, value=test_data)
            ws.cell(row=r, column=5, value=steps)
            ws.cell(row=r, column=6, value=expected)
            fit_row_to_text(ws, r)
            # The assumption may sit in any column (e.g. a max length assumed in
            # Precondition) - the whole case is then unconfirmed, but only the
            # expected result cell is highlighted.
            if any(needs_confirmation(v) for v in (purpose, precondition, test_data, steps, expected)):
                ws.cell(row=r, column=EXPECTED_COL).fill = CONFIRM_FILL
                highlighted.append(stt)
            stt += 1
            r += 1

    # Output must not freeze any row/column - clear panes on every sheet.
    for sheet in wb.worksheets:
        sheet.freeze_panes = None
    return wb, highlighted


def main():
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument('input', help='Path to the source .md file (viet-test-case output)')
    parser.add_argument('output', nargs='?', help='Path to the output .xlsx file (default: alongside input, same name)')
    parser.add_argument('--qc-name', default='Tuyền Võ', help='Ten QC dien vao muc "QC thuc hien" (mac dinh: Tuyền Võ)')
    args = parser.parse_args()

    input_path = Path(args.input)
    if not input_path.exists():
        print(f'Input file not found: {input_path}', file=sys.stderr)
        sys.exit(1)
    if not TEMPLATE_PATH.exists():
        print(f'Template file not found: {TEMPLATE_PATH}', file=sys.stderr)
        sys.exit(1)

    output_path = Path(args.output) if args.output else input_path.with_suffix('.xlsx')

    md_text = input_path.read_text(encoding='utf-8')
    tables = find_tables(md_text)
    if not tables:
        print('Khong tim thay bang test case (dang Markdown table) trong file input.', file=sys.stderr)
        sys.exit(1)

    table = tables[0]
    feature_name, redmine_url = extract_feature_info(md_text, input_path)

    try:
        wb, highlighted = build_workbook(table['header'], table['rows'], feature_name, redmine_url, args.qc_name)
    except ValueError as e:
        print(str(e), file=sys.stderr)
        sys.exit(1)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
    print(f'Da tao file Excel: {output_path}')
    if highlighted:
        print(f'Highlight {len(highlighted)} o "Ket qua mong muon" can confirm PO/BA - STT: '
              + ', '.join(str(s) for s in highlighted))
    else:
        print('Khong co o "Ket qua mong muon" nao can confirm PO/BA.')


if __name__ == '__main__':
    main()
