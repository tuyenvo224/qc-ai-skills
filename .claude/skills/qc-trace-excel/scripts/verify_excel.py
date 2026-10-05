"""Check an exported workbook against references/requirements.md. Prints PASS/FAIL per rule; exit 1 on any FAIL.

Usage: python verify_excel.py <workbook.xlsx> <OUT report folder>
"""
import csv
import json
import re
import sys
from pathlib import Path

import openpyxl

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

STATUSES = {"✅ Test được", "⚠ Test được, có giả định", "⏸ Chờ PO chốt", "⛔ Chưa test được"}
CASE_RE = re.compile(r"^(TC|SEC|E2E)-\d+$")


def read_csv(p):
    text = Path(p).read_text(encoding="utf-8-sig")
    first = text.splitlines()[0] if text else ""
    return list(csv.DictReader(text.splitlines(), delimiter=";" if first.count(";") >= first.count(",") else ","))


def main():
    xlsx, out = Path(sys.argv[1]), Path(sys.argv[2])
    wb = openpyxl.load_workbook(xlsx)
    data = json.loads((out / "data.json").read_text(encoding="utf-8"))
    qc = read_csv(out / "qc-cases.csv")
    live = [r["ID"] for r in qc if r["ID"][:2] != "TF" and "withdraw" not in (r.get("Trạng thái") or "").lower()]
    flows = [r["ID"] for r in qc if r["ID"].startswith("TF")]
    results = []

    def check(rule, ok, detail=""):
        results.append((rule, bool(ok), detail))

    # E2 sheet order / active
    names = wb.sheetnames
    check("E2 thứ tự sheet", names == ["Ghi chú", "TCs", "Ma trận FRD", "Test flow", "Tóm tắt"], " → ".join(names))
    check("E2 mở vào Ghi chú", wb.active.title == "Ghi chú", wb.active.title)

    # E4 row heights
    tall = [(ws.title, r) for ws in wb.worksheets for r, d in ws.row_dimensions.items() if (d.height or 0) > 409]
    check("E4 không dòng nào > 409pt", not tall, str(tall[:5]))

    ws = wb["TCs"]
    # T1 header
    check("T1 header tên chức năng / link FRD / Redmine / phiên bản",
          all(ws[c].value for c in ("F2", "F3", "F4", "F5")), " | ".join(str(ws[c].value)[:40] for c in ("F2", "F3", "F4", "F5")))
    # T3/T4 headings
    want = ["STT trường hợp kiểm thử", "Mục đích kiểm thử", "Precondition", "Test data", "Các bước thực hiện",
            "Kết quả mong muốn", "Kết quả hiện tại", "Mã lỗi", "Ghi chú", "QC thực hiện",
            "Có test được?", "Yêu cầu FRD liên quan", "Nguồn (file:dòng)"]
    got = [ws.cell(10, c).value for c in range(1, 14)]
    check("T3/T4 tiêu đề cột A–M", got == want, str([g for g, w in zip(got, want) if g != w]))
    dvs = [(d.formula1, str(d.sqref)) for d in ws.data_validations.dataValidation]
    check("T3 dropdown P/F/PE ở cột G", any(f == '"P,F,PE"' and s.startswith("G12") for f, s in dvs), str(dvs))

    # walk rows
    ids, groups, cont, empty, bad_status, first_row = [], [], 0, [], [], {}
    for r in range(12, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if not a:
            continue
        a = str(a)
        if CASE_RE.match(a):
            ids.append(a)
            first_row[a] = r
            if not ws.cell(r, 5).value or not ws.cell(r, 6).value:
                empty.append(a)
            if ws.cell(r, 11).value not in STATUSES:
                bad_status.append(a)
        elif "(tiếp" in a:
            cont += 1
        elif re.match(r"^[IVXLC]+\. ", a):
            groups.append(a)
    check("T2 đủ case TC/SEC/E2E (bỏ case rút)", sorted(ids) == sorted(live),
          f"Excel {len(ids)} · qc-cases {len(live)} · thiếu {sorted(set(live) - set(ids))[:8]} · thừa {sorted(set(ids) - set(live))[:8]}")
    check("T2 không có TF trong TCs", not any(i.startswith("TF") for i in ids))
    check("T6 cột K chỉ có 4 giá trị", not bad_status, str(bad_status[:8]))
    check("T10 không trống Các bước / Kết quả", not empty, str(empty[:8]))
    lane_words = ("Test ngay", "Chờ PO chốt", "Test + ghi thiếu", "— Chưa có")
    check("T5 tiêu đề nhóm không ghi làn", not any(w in g for g in groups for w in lane_words), str(groups[:3]))
    topic_names = [g["name"] for g in data["groups"]]
    shown = [re.sub(r"^[IVXLC]+\. ", "", g).rsplit(" (", 1)[0] for g in groups]
    in_order = [n for n in topic_names if n in shown]
    check("T5 nhóm theo Nhóm chức năng, đúng thứ tự FRD", shown[:len(in_order)] == in_order and
          all(s in topic_names or s.startswith("Tính năng ngoài FRD") for s in shown), " | ".join(shown[:6]))
    check("T9 sheet TCs không freeze", ws.freeze_panes is None, str(ws.freeze_panes))
    check("E4 case dài được tách dòng (tiếp)", True, f"{cont} dòng tiếp")

    # Ma trận
    mx = wb["Ma trận FRD"]
    mh = [mx.cell(1, c).value for c in range(1, 16)]
    want_m = ["R", "Nhóm chức năng", "Mã yêu cầu", "Mục FRD", "Loại", "Trích nguyên văn FRD", "BA", "BA IDs", "Chênh lệch BA",
              "TA", "TA IDs", "Chênh lệch TA", "Case QC", "Case test được ngay", "Tình trạng test"]
    check("X2 tiêu đề cột ma trận", mh == want_m, str([h for h, w in zip(mh, want_m) if h != w]))
    rows = data["rows"]
    rs = [mx.cell(r, 1).value for r in range(2, mx.max_row + 1) if mx.cell(r, 1).value]
    check("X1 đủ mọi yêu cầu FRD", rs == [x["r"] for x in rows], f"Excel {len(rs)} · data {len(rows)}")
    check("X2 mọi yêu cầu có Nhóm chức năng + Mã yêu cầu",
          all(mx.cell(r, 2).value and mx.cell(r, 3).value for r in range(2, len(rs) + 2)))
    bad_link = []
    for r in range(2, len(rs) + 2):
        h = mx.cell(r, 13).hyperlink
        if h and h.location:
            m = re.search(r"A(\d+)$", h.location)
            target = str(ws.cell(int(m.group(1)), 1).value) if m else ""
            listed = (mx.cell(r, 13).value or "").split()
            if target not in listed:
                bad_link.append(rs[r - 2])
    check("X3 link Case QC trỏ đúng case", not bad_link, str(bad_link[:8]))
    check("X4 bộ lọc + freeze tiêu đề", bool(mx.auto_filter.ref) and mx.freeze_panes == "B2", f"{mx.auto_filter.ref} {mx.freeze_panes}")
    known = set(ids) | set(flows) | {r["ID"] for r in qc}
    dangling = sorted({q for r in range(2, len(rs) + 2) for q in (mx.cell(r, 13).value or "").split() if q not in known})
    check("X3 mọi mã case trong ma trận đều tồn tại", not dangling, str(dangling[:8]))

    # Test flow / Tóm tắt
    tf = wb["Test flow"]
    tf_ids = [tf.cell(r, 1).value for r in range(2, tf.max_row + 1) if re.match(r"^TF-\d+$", str(tf.cell(r, 1).value or ""))]
    check("F1 đủ mọi TF", sorted(tf_ids) == sorted(flows), f"Excel {len(tf_ids)} · qc-cases {len(flows)}")
    sm = wb["Tóm tắt"]
    stext = " ".join(str(sm.cell(r, c).value or "") for r in range(1, sm.max_row + 1) for c in (1, 2))
    check("S1 Tóm tắt có số liệu cột K + câu hỏi PO", all(s in stext for s in STATUSES) and "CÂU HỎI CẦN PO CHỐT" in stext)

    # Ghi chú
    g = wb["Ghi chú"]
    gtext = "\n".join(f"{g.cell(r, 1).value or ''} | {g.cell(r, 2).value or ''}" for r in range(1, g.max_row + 1))
    for rule, needle in [
        ("G1 giới thiệu + nguồn", "File này là gì"), ("G2 vai trò sheet", "1. CÁC SHEET"),
        ("G3 định nghĩa yêu cầu FRD", "2. YÊU CẦU FRD (R-xxx) LÀ GÌ"), ("G3 ví dụ thật", "Ví dụ — mục"),
        ("G4 cột TCs", "3. SHEET TCs"), ("G4 cột Ma trận", "4. SHEET MA TRẬN FRD"), ("G5 trạng thái", "5. GIÁ TRỊ TRẠNG THÁI"),
        ("G6 cách đọc đúng", "6. CÁCH ĐỌC ĐÚNG"), ("G6 yêu cầu chưa có case", "Yêu cầu FRD nào chưa có case?"),
        ("G6 case kiểm yêu cầu nào", "Một case kiểm những yêu cầu nào?"), ("G6 nhiều–nhiều", "nhiều–nhiều"),
        ("G6 không gộp trùng", "không gộp các yêu cầu trùng ý"), ("G6 case ngoài FRD", "Case không gắn yêu cầu FRD nào?"),
        ("G6 dòng tiếp", "(tiếp 2)"), ("G6 trước khi log bug", "Trước khi log bug"),
        ("G7 bảng tra mã", "7. BẢNG TRA MÃ"), ("G8 lưu ý", "8. LƯU Ý KHI DÙNG"), ("G8 giới hạn", "Giới hạn")]:
        check(rule, needle in gtext)
    check("G6 số liệu khớp dữ liệu", f"{sum(1 for x in rows if not x['qc'].split())} yêu cầu" in gtext)
    check("E7 không dùng 'dòng FRD'", "dòng FRD" not in gtext)

    fails = [x for x in results if not x[1]]
    for rule, ok, detail in results:
        print(f"{'PASS' if ok else 'FAIL'}  {rule}" + (f"  — {detail}" if detail and (not ok or rule.startswith('E4 case')) else ""))
    print(f"\n{len(results) - len(fails)}/{len(results)} PASS")
    sys.exit(1 if fails else 0)


if __name__ == "__main__":
    main()
