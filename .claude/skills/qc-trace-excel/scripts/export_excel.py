"""Export a qc-frd-trace report folder to an Excel workbook built on the team's TC template.

Sheets: TCs (template layout, one row per case, grouped by feature group) · Ma trận FRD · Test flow · Tóm tắt
Usage:
  python export_excel.py <OUT> [--qc-dir <quality-control/.../feature>] [--ta-dir <technical-design/.../feature>]
                         [--template <Template_TCs.xlsx>] [--qc-name "Tên QC"] [--xlsx <output path>]
"""
import argparse
import copy
import csv
import datetime as dt
import json
import re
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.hyperlink import Hyperlink

if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")

DEFAULT_TEMPLATE = str(Path(__file__).resolve().parent.parent / "assets" / "Template_TCs.xlsx")  # ships with the skill
MAXCELL = 32000
OK, WARN, WAIT, NO = "✅ Test được", "⚠ Test được, có giả định", "⏸ Chờ PO chốt", "⛔ Chưa test được"
FILL = {OK: "FFE6F4EC", WARN: "FFFFF6DB", WAIT: "FFE9F0F6", NO: "FFFDECEA"}
ST_FILL = {"Đủ": "FFE6F4EC", "Một phần": "FFFFF6DB", "Lệch": "FFFDECEA", "Thiếu": "FFFFE3CC", "N/A": "FFEDEFF2"}
ROMAN = ["I", "II", "III", "IV", "V", "VI", "VII", "VIII", "IX", "X", "XI", "XII", "XIII", "XIV", "XV", "XVI", "XVII",
         "XVIII", "XIX", "XX", "XXI", "XXII", "XXIII", "XXIV", "XXV", "XXVI", "XXVII", "XXVIII", "XXIX", "XXX", "XXXI",
         "XXXII", "XXXIII", "XXXIV", "XXXV", "XXXVI", "XXXVII", "XXXVIII", "XXXIX", "XL", "XLI", "XLII", "XLIII",
         "XLIV", "XLV", "XLVI", "XLVII", "XLVIII", "XLIX", "L", "LI", "LII", "LIII", "LIV", "LV", "LVI", "LVII", "LVIII"]
IDRE = re.compile(r"(?:BR|VR|AC|UC|AF)-\d+")


# ---------------------------------------------------------------- reading helpers
def read_csv(path):
    if not Path(path).exists():
        return []
    text = Path(path).read_text(encoding="utf-8-sig")
    first = text.splitlines()[0] if text else ""
    delim = ";" if first.count(";") >= first.count(",") else ","
    return list(csv.DictReader(text.splitlines(), delimiter=delim))


def col(row, *names):
    for n in names:
        for k in row:
            if k and k.strip().startswith(n):
                return (row[k] or "").strip()
    return ""


def cells(line):
    parts = re.split(r"(?<!\\)\|", line.strip())[1:-1]
    return [p.strip().replace("\\|", "|") for p in parts]


def clean(s):
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s or "")
    s = re.sub(r"`([^`]*)`", r"\1", s)
    s = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", s)
    return s.strip()


def table_text(rows):
    """markdown table rows (list of cell lists, first = header) -> readable lines."""
    if not rows:
        return ""
    head, body = rows[0], rows[1:]
    out = []
    if len(head) == 2 and not head[0]:            # key/value meta table
        return "\n".join(f"{clean(r[0])}: {clean(r[1])}" for r in body if len(r) > 1)
    for r in body:
        r = r + [""] * (len(head) - len(r))
        if len(head) == 2:
            out.append(f"• {clean(r[0])} → {clean(r[1])}")
        else:
            out.append("• " + " · ".join(f"{clean(h)}: {clean(v)}" for h, v in zip(head, r) if clean(v)))
    return "\n".join(out)


def parse_block(lines):
    """Split a case block into meta table, labelled sections, bare tables and prose."""
    meta, sections, bare, prose = {}, [], [], []
    cur, i = None, 0
    while i < len(lines):
        ln = lines[i]
        if ln.startswith("|"):
            tbl = []
            while i < len(lines) and lines[i].startswith("|"):
                if not re.match(r"^\|\s*:?-{2,}", lines[i]):
                    tbl.append(cells(lines[i]))
                i += 1
            if not meta and tbl and len(tbl[0]) == 2 and not tbl[0][0] and cur is None:
                for r in tbl[1:]:
                    if len(r) > 1:
                        meta[clean(r[0])] = r[1]
            elif cur is not None:
                cur[1].append(("table", tbl))
            else:
                bare.append(tbl)
            continue
        m = re.match(r"^\*\*([^*]+?)\*\*:?\s*(.*)$", ln.strip())
        if m and len(m.group(1)) < 60:
            cur = (m.group(1).strip().rstrip(":"), [])
            sections.append(cur)
            if m.group(2).strip():
                cur[1].append(("text", m.group(2).strip()))
        elif ln.strip() in ("---", ""):
            if ln.strip() == "---":
                cur = None
        elif re.match(r"^\s*(\d+\.|[-*])\s", ln) and cur is not None:
            cur[1].append(("text", ln.strip()))
        else:
            (cur[1] if cur is not None and cur[0].lower().startswith(("steps", "attack")) else prose).append(
                ("text", ln.strip()) if cur is not None and cur[0].lower().startswith(("steps", "attack")) else ln.strip())
        i += 1
    return meta, sections, bare, [p for p in prose if isinstance(p, str)]


def sect_text(items):
    out = []
    for kind, v in items:
        out.append(table_text(v) if kind == "table" else clean(v))
    return "\n".join(x for x in out if x)


def case_blocks(path):
    """{case_id: (title, lines)} for every heading that starts with `ID`."""
    lines = Path(path).read_text(encoding="utf-8").splitlines()
    heads = [(i, len(m.group(1)), m.group(2), m.group(3)) for i, l in enumerate(lines)
             for m in [re.match(r"^(#{2,5})\s+`?~*([A-Z][A-Z0-9]{0,4}-\d+)~*`?~*\s*[—-]*\s*(.*)$", l)] if m]
    out = {}
    for k, (i, lvl, cid, title) in enumerate(heads):
        end = len(lines)
        for j, lvl2, _, _ in heads[k + 1:]:
            if lvl2 <= lvl:
                end = j
                break
        for j in range(i + 1, end):
            if re.match(r"^#{1,%d}\s" % lvl, lines[j]):
                end = j
                break
        out.setdefault(cid, (clean(title), lines[i + 1:end], i + 1))
    return out


def autofit(ws, first_row, last_row, max_col):
    """Excel does not grow wrapped rows on open, so compute each row height from its text."""
    import math
    too_tall = []
    merged = {(r, c) for rng in ws.merged_cells.ranges for r in range(rng.min_row, rng.max_row + 1)
              for c in range(rng.min_col, rng.max_col + 1) if not (r == rng.min_row and c == rng.min_col)}
    for r in range(first_row, last_row + 1):
        lines_max = 1
        for c in range(1, max_col + 1):
            if (r, c) in merged:
                continue
            v = ws.cell(r, c).value
            if v is None or v == "":
                continue
            letter = openpyxl.utils.get_column_letter(c)
            width = ws.column_dimensions[letter].width or 10
            chars = max(int(width * 1.05) - 1, 5)
            n = sum(max(1, math.ceil(len(part) / chars)) for part in str(v).split(chr(10)))
            lines_max = max(lines_max, n)
        size = ws.cell(r, 1).font.sz or 11
        h = lines_max * size * 1.35 + 4
        if h > 409:
            too_tall.append(r)
        ws.row_dimensions[r].height = min(409, h)
    return too_tall


MAX_LINES = 25  # ~409pt Excel row limit at 11pt


def split_cells(vals, widths):
    """Split long cell texts into row chunks so no row exceeds Excel's max height."""
    import math
    per_col = []
    for v, w in zip(vals, widths):
        if not isinstance(v, str) or not v:
            per_col.append([v])
            continue
        chars = max(int((w or 10) * 1.05) - 1, 5)
        chunks, cur, used = [], [], 0
        for para in v.split(chr(10)):
            pieces = [para[i:i + chars * MAX_LINES] for i in range(0, max(len(para), 1), chars * MAX_LINES)] or [""]
            for piece in pieces:
                n = max(1, math.ceil(len(piece) / chars))
                if cur and used + n > MAX_LINES:
                    chunks.append(chr(10).join(cur))
                    cur, used = [], 0
                cur.append(piece)
                used += n
        chunks.append(chr(10).join(cur))
        per_col.append(chunks)
    n = max(len(c) for c in per_col)
    return [[(c[k] if k < len(c) else None) for c in per_col] for k in range(n)]



def pick_example(rows, out):
    """A short FRD list section (5–10 short rows) to illustrate what one requirement is; plus 0-row sections."""
    by = {}
    for r in rows:
        by.setdefault(r["sec"], []).append(r)
    ex, ex_sec = [], ""
    for sec, rs in by.items():
        if 5 <= len(rs) <= 10 and all(len(r["q"]) <= 80 for r in rs):
            ex, ex_sec = [(r["r"], r["q"], r["typ"]) for r in rs], sec
            break
    empty = ""
    cs = out / "coverage-sections.json"
    if cs.exists():
        e = json.loads(cs.read_text(encoding="utf-8")).get("empty", [])
        empty = ", ".join(f"§{x['num']} {x['title']}" for x in e)
    return dict(example=ex, example_sec=ex_sec, empty_sections=empty)


# ---------------------------------------------------------------- main
def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("out")
    ap.add_argument("--qc-dir", help="default: run-meta.json → qc_dir")
    ap.add_argument("--ta-dir")
    ap.add_argument("--template", default=DEFAULT_TEMPLATE)
    ap.add_argument("--qc-name")
    ap.add_argument("--xlsx")
    a = ap.parse_args()
    out = Path(a.out)
    rmeta = json.loads((out / "run-meta.json").read_text(encoding="utf-8")) if (out / "run-meta.json").exists() else {}
    a.qc_dir = a.qc_dir or rmeta.get("qc_dir")
    a.ta_dir = a.ta_dir or rmeta.get("ta_dir")
    if not a.qc_dir or not Path(a.qc_dir).is_dir():
        sys.exit("STOP: QC folder unknown or missing — pass --qc-dir, or add qc_dir to run-meta.json in the report folder.")
    if not Path(a.template).exists():
        sys.exit(f"STOP: template not found: {a.template} — pass --template.")
    qcdir = Path(a.qc_dir)

    data = json.loads((out / "data.json").read_text(encoding="utf-8"))
    narr = json.loads((out / "narrative.json").read_text(encoding="utf-8")) if (out / "narrative.json").exists() else {}
    meta = data.get("meta", {})
    qc_rows = read_csv(out / "qc-cases.csv")
    extra = {col(x, "BA ID"): col(x, "Loại") for x in read_csv(out / "ba-extra.csv")}

    # assumptions A-* from qc-prereq.md
    assume = {}
    pq = out / "qc-prereq.md"
    if pq.exists():
        for ln in pq.read_text(encoding="utf-8").splitlines():
            c = cells(ln) if ln.startswith("|") else []
            if c and re.fullmatch(r"`?A-\d+`?", c[0]):
                assume[c[0].strip("`")] = clean(c[1]) if len(c) > 1 else ""

    # element names from ui-contract
    el_name = {}
    if a.ta_dir and (Path(a.ta_dir) / "ui-contract.md").exists():
        for ln in (Path(a.ta_dir) / "ui-contract.md").read_text(encoding="utf-8").splitlines():
            c = cells(ln) if ln.startswith("|") else []
            if len(c) >= 4 and re.fullmatch(r"`?EL-\d+`?", c[0]):
                el_name.setdefault(c[0].strip("`"), clean(c[3]))

    def enrich(s):
        seen = set()

        def rep(m):
            e = m.group(0)
            if e in seen or e not in el_name:
                return e
            seen.add(e)
            return f"{e} ({el_name[e]})"
        return re.sub(r"EL-\d+", rep, s)

    # blocks from every QC file
    blocks = {}
    for f in sorted(qcdir.glob("*.md")):
        for cid, b in case_blocks(f).items():
            blocks.setdefault(cid, (f.name,) + b)

    # R rows & groups
    rows = data["rows"]
    r_by_ba = {}
    for r in rows:
        for i in r["ba"].split():
            r_by_ba.setdefault(i, []).append(r["r"])
    groups = data["groups"]
    row_by_r = {r["r"]: r for r in rows}
    g_of_r = {}
    for gi, g in enumerate(groups):
        for r in g["rs"]:
            g_of_r.setdefault(r, gi)

    # ---------- build case records
    cases, flows, skipped = [], [], []
    for q in qc_rows:
        cid = col(q, "ID")
        st = col(q, "Trạng thái")
        fname, title, blines, ln0 = blocks.get(cid, (col(q, "File"), col(q, "Tiêu đề"), [], col(q, "Dòng")))
        m, secs, bare, prose = parse_block(blines)
        if cid.startswith("TF"):
            steps = ""
            for t in bare:
                steps += table_text(t) + "\n"
            flows.append(dict(id=cid, title=title or col(q, "Tiêu đề"), cls=clean(m.get("Scenario class", "")),
                              setup=clean(m.get("Setup", "")), traces=clean(m.get("Traces", "")),
                              steps=enrich(steps.strip() or sect_text(sum((s[1] for s in secs), []))),
                              src=f"{fname}:{ln0}"))
            continue
        if "withdraw" in st.lower():
            skipped.append((cid, title, st))
            continue
        ba_ids = sorted(set(IDRE.findall(col(q, "Traces") + " " + m.get("Traces", ""))))
        rs = sorted({r for i in ba_ids for r in r_by_ba.get(i, [])}, key=lambda x: int(x[2:]))

        def pick(*keys):
            return [s for s in secs if s[0].lower().startswith(keys)]
        pre = "\n".join(clean(m[k]) for k in m if k.lower().startswith(("precondition", "setup", "actor", "điểm điều khiển")))
        data_s = sect_text(sum((s[1] for s in pick("input") if "expected" not in s[0].lower()), []))
        steps = sect_text(sum((s[1] for s in pick("steps", "attack")), []))
        exp = sect_text(sum((s[1] for s in pick("expected", "cross-layer", "input data / expected")), []))
        if bare and not steps:  # E2E journey table: action / system result
            t = bare[0]
            steps = "\n".join(f"{clean(r[0])}. {clean(r[1])}" for r in t[1:] if len(r) > 1)
            exp_j = "\n".join(f"{clean(r[0])}. {clean(r[2])}" for r in t[1:] if len(r) > 2)
            exp = (exp_j + ("\n\nKiểm chéo các tầng:\n" + exp if exp else "")).strip()
        if bare and not data_s and cid.startswith("E2E"):
            data_s = "Dữ liệu nằm ngay trong từng bước (cột Các bước thực hiện)."
        if pick("input data / expected") and not data_s:
            data_s = "Xem bảng trong cột Kết quả mong muốn (dữ liệu và kết quả nằm chung một bảng)."
        note_prose = clean(" ".join(prose))[:600]

        # testability
        deps = col(q, "Phụ thuộc")
        a_ids = re.findall(r"A-\d+", deps)
        reasons = []
        status = OK
        g_idx = [g_of_r[r] for r in rs if r in g_of_r]
        impl = {row_by_r[r]["impl"] for r in rs if row_by_r.get(r, {}).get("impl")}
        against = [i for i in ba_ids if extra.get(i) == "Trái FRD"]
        if "BLOCK" in st.upper():
            status, reasons = NO, [st if st.upper().startswith("BLOCK") else f"BLOCKED: {st}"]
        elif "AUTO" in st.upper():
            status, reasons = NO, ["Chỉ chạy được bằng test tự động của dev (làn AUTO), không có đường chạy tay"]
        elif impl and impl <= {"Chưa có"}:
            status, reasons = NO, ["Tính năng liên quan chưa có trong code"]
        elif against or (impl and impl <= {"Khác FRD", "Chưa có"}):
            status = WAIT
            reasons = ["Kiểm hành vi BA/TA làm khác FRD" + (f" ({', '.join(against[:5])} trái FRD)" if against else "")]
        elif a_ids or "OQ" in st:
            status = WARN
        if a_ids:
            reasons += [f"Giả định {x}: {assume.get(x, '')}".strip(": ") for x in a_ids]
        if "OQ" in st and status != NO:
            reasons.append(st)
        gi = max(set(g_idx), key=g_idx.count) if g_idx else None
        cases.append(dict(id=cid, title=title, ba=ba_ids, rs=rs, group=gi, status=status,
                          purpose=f"{title}\n(Phủ: {', '.join(ba_ids[:12])}{' …' if len(ba_ids) > 12 else ''})" if ba_ids else title,
                          pre=enrich(pre), data=enrich(data_s), steps=enrich(steps), exp=enrich(exp),
                          note="\n".join(reasons + ([note_prose] if note_prose else [])),
                          layer=col(q, "Layer"), src=f"{fname}:{ln0}"))

    # ---------- workbook
    wb = openpyxl.load_workbook(a.template)
    ws = wb.worksheets[0]
    ws.title = "TCs"
    tpl_group = {c.column: copy.copy(c._style) for c in ws[12]}
    tpl_data = {c.column: copy.copy(c._style) for c in ws[13]}
    for rng in list(ws.merged_cells.ranges):
        if rng.min_row >= 11:
            ws.unmerge_cells(str(rng))
    for r in ws.iter_rows(min_row=11, max_row=ws.max_row):
        for c in r:
            c.value = None
    feat = narr.get("feature") or meta.get("name", "")
    rm = data.get("redmine", [])
    ws["F2"] = f"{feat} ({', '.join('#' + str(x['id']) for x in rm if x.get('tracker') == 'Story') or ''})".replace(" ()", "")
    ws["F3"] = meta.get("url", "")
    if meta.get("url"):
        ws["F3"].hyperlink = meta["url"]
    ws["F4"] = " · ".join(f"https://redmine.gotit.vn/issues/{x['id']}" for x in rm if x.get("tracker") == "Story")
    if a.qc_name:
        ws["H3"] = a.qc_name
    ws["I3"] = dt.datetime.now().replace(hour=0, minute=0, second=0, microsecond=0)
    ver = ""
    fr = out / "frd-reqs.md"
    if fr.exists():
        sec = False
        for ln in fr.read_text(encoding="utf-8").splitlines():
            if ln.startswith("## "):
                sec = ln.lower().startswith("## lịch sử")
            elif sec and re.match(r"^\|\s*v\d", ln):
                c = cells(ln)
                ver = f"{c[0]} · {c[1]}"
    ws["F5"] = f"FRD {ver or ''} (BookStack rev {meta.get('revision', '')}, cập nhật {str(meta.get('updated_at', ''))[:10]})" if meta else ""
    for col_letter, head in (("K", "Có test được?"), ("L", "Yêu cầu FRD liên quan"), ("M", "Nguồn (file:dòng)")):
        ws[f"{col_letter}10"] = head
        ws[f"{col_letter}10"]._style = copy.copy(ws["J10"]._style)
    ws.column_dimensions["K"].width = 30
    ws.column_dimensions["L"].width = 28
    ws.column_dimensions["M"].width = 26
    for cl, w in (("B", 45), ("C", 42), ("D", 42), ("E", 75), ("F", 85), ("I", 50)):
        ws.column_dimensions[cl].width = max(ws.column_dimensions[cl].width or 0, w)

    # topic groups ('Nhóm chức năng') keep FRD catalog order; legacy code groups are sorted by lane
    order = list(range(len(groups))) if groups and groups[0].get("key") else         sorted(range(len(groups)), key=lambda i: (groups[i]["lane"], groups[i]["name"]))
    buckets = {i: [] for i in order}
    buckets[None] = []
    for c in cases:
        buckets.setdefault(c["group"], []).append(c)
    lane_name = {1: "Test ngay", 2: "Test + ghi thiếu", 3: "Chờ PO chốt", 4: "Chưa có"}
    row, gno, first_row = 12, 0, {}
    wrap = Alignment(wrap_text=True, vertical="top")
    for gi in order + [None]:
        items = buckets.get(gi) or []
        if not items:
            continue
        name = groups[gi]["name"] if gi is not None else "Tính năng ngoài FRD (BA tự thêm)"
        lane = "" if gi is None or groups[gi].get("key") else f" — {lane_name.get(groups[gi]['lane'], '')}"
        for cc in range(1, 14):
            ws.cell(row, cc)._style = copy.copy(tpl_group.get(cc, tpl_group.get(1)))
        ws.cell(row, 1, f"{ROMAN[gno] if gno < len(ROMAN) else gno + 1}. {name}{lane} ({len(items)} case)")
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        row += 1
        gno += 1
        for c in sorted(items, key=lambda x: (x["id"].split("-")[0] != "TC", x["id"])):
            vals = [c["id"], c["purpose"], c["pre"], c["data"], c["steps"], c["exp"], None, None, c["note"], None,
                    c["status"], (", ".join(c["rs"][:15]) + (f" … (+{len(c['rs']) - 15})" if len(c["rs"]) > 15 else "")) if c["rs"]
                    else ("Không có yêu cầu FRD tương ứng (mã BA tự thêm ngoài FRD)" if c["ba"] else "—"), c["src"]]
            widths = [ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width for i in range(1, 14)]
            parts = split_cells(vals, widths)
            first_row[c["id"]] = row
            for k, part in enumerate(parts):
                if k:
                    part[0] = f"{c['id']} (tiếp {k + 1})"
                    for keep in (10, 11, 12):  # status / FRD refs / source only on the first row
                        part[keep] = None
                for cc, v in enumerate(part, 1):
                    cell = ws.cell(row, cc)
                    cell._style = copy.copy(tpl_data.get(cc, tpl_data.get(1)))
                    cell.value = (v[:MAXCELL] if isinstance(v, str) else v)
                    cell.alignment = wrap
                ws.cell(row, 11).fill = PatternFill("solid", fgColor=FILL[c["status"]])
                row += 1
    last = row - 1
    ws.data_validations.dataValidation = []
    dv = DataValidation(type="list", formula1='"P,F,PE"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"G12:G{max(last, 12)}")
    ws.freeze_panes = None
    tall = autofit(ws, 12, last, 13)
    ws.auto_filter.ref = f"A10:M{last}"

    # ---------- Ma trận FRD
    hdr_style = copy.copy(ws["A10"]._style)
    mx = wb.create_sheet("Ma trận FRD")
    heads = ["R", "Nhóm chức năng", "Mã yêu cầu", "Mục FRD", "Loại", "Trích nguyên văn FRD", "BA", "BA IDs", "Chênh lệch BA", "TA", "TA IDs",
             "Chênh lệch TA", "Case QC", "Case test được ngay", "Tình trạng test"]
    widths = [8, 26, 26, 22, 11, 60, 10, 18, 45, 10, 18, 45, 28, 22, 26]
    for i, (h, w) in enumerate(zip(heads, widths), 1):
        c = mx.cell(1, i, h)
        c._style = copy.copy(hdr_style)
        mx.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    status_of = {c["id"]: c["status"] for c in cases}
    for k, r in enumerate(rows, 2):
        qcs = r["qc"].split()
        live = [q for q in qcs if q in status_of]
        ok_now = [q for q in live if status_of[q] in (OK, WARN)]
        if not qcs:
            tinh = "Không có case QC"
        elif ok_now:
            tinh = f"Có {len(ok_now)} case test được"
        else:
            tinh = "Chỉ có case chưa test được / chờ PO"
        vals = [r["r"], r.get("topic_name", ""), r.get("grp", ""), r["sec"], r["typ"], r["q"], r["st"], r["ba"], r["diff"], r.get("ta_st", ""), r.get("ta", ""),
                r.get("ta_diff", ""), " ".join(qcs), " ".join(ok_now), tinh]
        for i, v in enumerate(vals, 1):
            c = mx.cell(k, i, v)
            c.alignment = wrap
        for ci, key in ((7, r["st"]), (10, r.get("ta_st", ""))):
            if key in ST_FILL:
                mx.cell(k, ci).fill = PatternFill("solid", fgColor=ST_FILL[key])
        target = next((first_row[q] for q in qcs if q in first_row), None)
        if target:
            mx.cell(k, 13).hyperlink = Hyperlink(ref=f"M{k}", location=f"'TCs'!A{target}", display=" ".join(qcs))
            mx.cell(k, 13).font = Font(color="FF1F5F8B", underline="single")
        if tinh.startswith("Không"):
            mx.cell(k, 15).fill = PatternFill("solid", fgColor="FFFFE3CC")
    mx.freeze_panes = "B2"
    tall += autofit(mx, 2, len(rows) + 1, 15)
    mx.auto_filter.ref = f"A1:O{len(rows) + 1}"

    # ---------- Test flow
    tf = wb.create_sheet("Test flow")
    th = ["TF", "Tên luồng", "Loại", "Setup", "Các bước (Bước → Assert · Chạy case)", "Traces", "Nguồn"]
    for i, (h, w) in enumerate(zip(th, [9, 40, 12, 40, 90, 30, 22]), 1):
        c = tf.cell(1, i, h)
        c._style = copy.copy(hdr_style)
        tf.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
    k = 2
    tfw = [tf.column_dimensions[openpyxl.utils.get_column_letter(i)].width for i in range(1, 8)]
    for f in flows:
        for j, part in enumerate(split_cells([f["id"], f["title"], f["cls"], f["setup"], f["steps"][:MAXCELL], f["traces"], f["src"]], tfw)):
            if j:
                part[0] = f"{f['id']} (tiếp {j + 1})"
            for i, v in enumerate(part, 1):
                tf.cell(k, i, v).alignment = wrap
            k += 1
    tf.freeze_panes = "B2"
    tall += autofit(tf, 2, k - 1, 7)

    # ---------- Tóm tắt
    sm = wb.create_sheet("Tóm tắt")
    sm.column_dimensions["A"].width = 38
    sm.column_dimensions["B"].width = 90
    cnt = {s: sum(1 for c in cases if c["status"] == s) for s in (OK, WARN, WAIT, NO)}
    st = data["stats"]
    lines = [("TÓM TẮT", ""), ("Tính năng", feat), ("FRD", meta.get("url", "")), ("Ngày xuất", dt.date.today().isoformat()), ("", ""),
             ("SHEET TCs", f"{len(cases)} case (không tính {len(skipped)} case đã rút, {len(flows)} test flow nằm ở sheet Test flow)")]
    lines += [(s, f"{cnt[s]} case") for s in (OK, WARN, WAIT, NO)]
    lines += [("", ""), ("Ý NGHĨA CỘT 'CÓ TEST ĐƯỢC?'", ""),
              (OK, "Code đã có, case không vướng giả định hay điểm cần PO chốt"),
              (WARN, "Chạy được, nhưng expected dựa trên giả định A-* hoặc câu hỏi chưa chốt — ghi rõ ở cột Ghi chú"),
              (WAIT, "Case kiểm hành vi mà BA/TA/code làm khác FRD — chạy được, nhưng pass/fail phụ thuộc PO chọn FRD hay spec Redmine"),
              (NO, "BLOCKED, tính năng code chưa có, hoặc chỉ chạy được bằng test tự động của dev"),
              ("", ""), ("SHEET MA TRẬN FRD", f"{st['frd_total']} yêu cầu FRD. BA: " + " · ".join(f"{k} {v}" for k, v in st["frd_vs_ba"].items())
                                                    + ((" | TA: " + " · ".join(f"{k} {v}" for k, v in st.get("frd_vs_ta", {}).items())) if st.get("frd_vs_ta") else "")),
              ("Yêu cầu FRD không có case QC", str(sum(1 for r in rows if not r["qc"].split()))),
              ("Cách dùng", "Lọc cột 'Tình trạng test'. Bấm vào ô 'Case QC' để nhảy tới case đầu tiên ở sheet TCs."),
              ("", "")]
    if skipped:
        lines += [("CASE ĐÃ RÚT (không đưa vào TCs)", ", ".join(f"{s[0]}" for s in skipped))]
    if narr.get("po_questions"):
        lines += [("", ""), ("CÂU HỎI CẦN PO CHỐT", "")]
        lines += [(f"{i}.", clean(q.get("q", "")) + (" — " + clean(q.get("detail", "")) if q.get("detail") else ""))
                  for i, q in enumerate(narr["po_questions"], 1)]
    for k, (x, y) in enumerate(lines, 1):
        sm.cell(k, 1, x).font = Font(bold=bool(x and (x.isupper() or x in FILL)))
        sm.cell(k, 2, y).alignment = wrap
        if x in FILL:
            sm.cell(k, 1).fill = PatternFill("solid", fgColor=FILL[x])

    # ---------- Ghi chú (usage guide, first sheet)
    import collections
    sys.path.insert(0, str(Path(__file__).parent))
    import notes_sheet
    nocase = [r for r in rows if not r["qc"].split()]
    topic_cnt = collections.Counter(r.get("topic_name") or r["grp"].split(".")[0] for r in nocase if r.get("grp"))
    per_case = collections.Counter(q for r in rows for q in r["qc"].split())
    n_cont = sum(1 for rr in range(12, last + 1) if "(tiếp" in str(ws.cell(rr, 1).value or ""))
    lim = [x for x in [narr.get("footer", "")] if x]
    notes_sheet.build(wb, ctx=dict(
        feature=feat, date=dt.date.today().strftime("%d/%m/%Y"),
        sources=" · ".join(narr.get("sources", [])) or meta.get("url", ""),
        n_cases=len(cases), n_flows=len(flows), n_rows=len(rows), n_withdrawn=len(skipped),
        test=cnt, OK=OK, WARN=WARN, WAIT=WAIT, NO=NO,
        n_keys=len({r.get("grp") for r in rows if r.get("grp")}), n_topics=len({r["grp"].split(".")[0] for r in rows if r.get("grp")}),
        n_nocase=len(nocase), n_nocase_keys=len({r.get("grp") for r in nocase if r.get("grp")}), n_nocase_topics=len(topic_cnt),
        top_nocase_topics=", ".join(f"{k}: {v} yêu cầu" for k, v in topic_cnt.most_common(4)),
        **pick_example(rows, out),
        max_r_per_case=max(per_case.values()) if per_case else 0,
        n_rows_with_case=sum(1 for r in rows if r["qc"].split()), n_distinct_cases=len(per_case),
        n_case_no_frd=sum(1 for c in cases if not c["rs"]), n_cont=n_cont,
        ents=notes_sheet.ent_tables(a.ta_dir),
        extra_notes=narr.get("excel_notes", []), has_code=bool(data["stats"].get("impl")),
        limits=" ".join(lim) or "Đánh giá Đủ/Lệch/Thiếu và cột K do AI thực hiện; kiểm lại trước khi log bug."))
    wb.active = 0

    dest = Path(a.xlsx) if a.xlsx else out / f"{out.parent.name}_TCs.xlsx"
    if dest.exists():
        import shutil
        backup = dest.with_name(dest.stem + "_backup" + dest.suffix)
        try:
            shutil.copyfile(dest, backup)
        except PermissionError:
            sys.exit(f"STOP: {dest.name} is open in Excel — close it and run again.")
    try:
        wb.save(dest)
    except PermissionError:
        sys.exit(f"STOP: {dest.name} is open in Excel — close it and run again.")
    empty = [c["id"] for c in cases if not c["steps"] or not c["exp"]]
    print(f"OK {dest}")
    print(f"cases {len(cases)} (+{len(skipped)} withdrawn skipped) · flows {len(flows)} · matrix rows {len(rows)}")
    print("testability: " + " · ".join(f"{s} {cnt[s]}" for s in cnt))
    print(f"cases with empty steps/expected: {len(empty)} {' '.join(empty[:20])}")
    print(f"rows taller than Excel max (409pt, text may still be cut): {len(tall)}")
    missing = [col(q, 'ID') for q in qc_rows if col(q, 'ID') not in blocks and "withdraw" not in col(q, "Trạng thái").lower()]
    if missing:
        print(f"WARNING: {len(missing)} case ids not found in QC files: {' '.join(missing[:20])}")


if __name__ == "__main__":
    main()
