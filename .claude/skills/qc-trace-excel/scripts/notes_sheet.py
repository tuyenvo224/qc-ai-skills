"""Build the 'Ghi chú' sheet: how to use the workbook, column meanings, how to read it correctly.

All numbers are computed from the same data that fills the other sheets.
"""
import re
from pathlib import Path

from openpyxl.styles import Alignment, Font, PatternFill

HEAD_FILL = PatternFill("solid", fgColor="FF1F5F8B")
SUB_FILL = PatternFill("solid", fgColor="FFE3EEF6")
WRAP = Alignment(wrap_text=True, vertical="top")


def ent_tables(ta_dir):
    """ENT-xxx -> (entity, table) from TA analysis.md / design.md data-model tables."""
    out = {}
    if not ta_dir:
        return out
    for name in ("analysis.md", "design.md"):
        p = Path(ta_dir) / name
        if not p.exists():
            continue
        for ln in p.read_text(encoding="utf-8").splitlines():
            m = re.match(r"^\|\s*`?(ENT-\d+)`?\s*\|\s*`?([^|`]+)`?[^|]*\|\s*`([a-z_][a-z0-9_]*)`\s*\|", ln)
            if m:
                out.setdefault(m.group(1), (m.group(2).strip(), m.group(3)))
    return out


def build(wb, *, ctx):
    """ctx: dict with numbers and labels prepared by export_excel.main()."""
    ws = wb.create_sheet("Ghi chú", 0)
    ws.column_dimensions["A"].width = 34
    ws.column_dimensions["B"].width = 115
    ws.sheet_view.showGridLines = False
    r = [1]

    def title(text):
        c = ws.cell(r[0], 1, text)
        c.font = Font(bold=True, size=14, color="FF1F5F8B")
        r[0] += 2

    def section(text):
        for col in (1, 2):
            ws.cell(r[0], col).fill = HEAD_FILL
        ws.cell(r[0], 1, text).font = Font(bold=True, color="FFFFFFFF", size=12)
        ws.row_dimensions[r[0]].height = 20
        r[0] += 1

    def sub(text):
        for col in (1, 2):
            ws.cell(r[0], col).fill = SUB_FILL
        ws.cell(r[0], 1, text).font = Font(bold=True)
        r[0] += 1

    def row(a, b=""):
        ca, cb = ws.cell(r[0], 1, a), ws.cell(r[0], 2, b)
        ca.alignment = cb.alignment = WRAP
        ca.font = Font(bold=bool(a))
        r[0] += 1

    def gap():
        r[0] += 1

    c = ctx
    t = c["test"]
    title(f"HƯỚNG DẪN ĐỌC FILE — {c['feature']}")
    row("File này là gì", "Bộ test case của tính năng, kèm ma trận truy vết từng yêu cầu trong FRD (spec gốc của PO/BA) tới tài liệu BA, TA "
        "và test case QC do dev sinh bằng AI. Mục tiêu: biết phải test gì, case nào chạy được ngay, yêu cầu FRD nào chưa có case.")
    row("Nguồn dữ liệu", c["sources"])
    row("Ngày xuất", c["date"])
    row("Sinh bởi", "Skill qc-frd-trace (phân tích, ra report + artifact) rồi qc-trace-excel (xuất file này). "
        "Chạy lại 2 skill khi FRD, tài liệu dev hoặc code thay đổi.")
    gap()

    section("1. CÁC SHEET VÀ THỨ TỰ NÊN DÙNG")
    row("Ghi chú", "Sheet này — đọc trước.")
    row("TCs", f"Sheet làm việc chính, theo mẫu test case của team. {c['n_cases']} case, gom theo nhóm chức năng, "
        "nhóm dễ test nhất ở trên. Ghi kết quả chạy ở đây.")
    row("Ma trận FRD", f"{c['n_rows']} yêu cầu FRD (R-xxx), mỗi yêu cầu một dòng, cho biết tài liệu BA/TA có đúng FRD không và case QC nào kiểm nó. "
        "Dùng để tìm yêu cầu chưa có case, hoặc tra một yêu cầu cụ thể.")
    row("Test flow", f"{c['n_flows']} luồng kiểm thử (TF-*). Mỗi luồng là thứ tự chạy nhiều case TC nối nhau. Không trùng với sheet TCs.")
    row("Tóm tắt", "Số liệu tổng, danh sách case đã rút, câu hỏi cần PO chốt.")
    row("Thứ tự đề xuất", "① Đọc mục 2, 5 và 6 của sheet này → ② gửi PO các câu hỏi ở sheet Tóm tắt → ③ ở sheet TCs lọc cột K = "
        f"\"{c['OK']}\" và chạy trước → ④ ghi P/F/PE → ⑤ dùng sheet Ma trận để báo cáo yêu cầu FRD chưa có case.")
    gap()

    section("2. YÊU CẦU FRD (R-xxx) LÀ GÌ")
    row("Định nghĩa", "Mỗi mã R-xxx là MỘT câu trong FRD mà QC kiểm được — một quy tắc field, validation, message, chuyển trạng thái, "
        "nhánh theo loại, thông báo, tác động lên hệ thống ngoài, giới hạn hoặc hành vi UI. Trong file này \"yêu cầu FRD\" luôn có nghĩa đó.")
    row("Câu ghép thì tách", "Một câu FRD nói hai điều thì thành hai yêu cầu.")
    row("Không gộp, không diễn giải", "Giữ nguyên ý gốc, không tóm tắt, không thêm ý. Cột Trích nguyên văn chép đúng câu gốc (≤ 300 ký tự) "
        "và đã được script kiểm là có thật trong FRD.")
    row("Câu lặp vẫn là yêu cầu riêng", "FRD nhắc cùng một ý ở nhiều mục (mô tả, use case, luồng, yêu cầu chức năng, validation, AC) "
        "thì mỗi chỗ là một yêu cầu riêng, có chung Mã yêu cầu.")
    row("Thứ tự", f"Đánh số theo thứ tự trong FRD: R-001 là câu đầu, R-{c['n_rows']:03d} là câu cuối.")
    row("Phần không có yêu cầu", "Mục chỉ là bối cảnh, không có gì để test, thì không sinh yêu cầu nào"
        + (f" (ở tính năng này: {c['empty_sections']})." if c.get("empty_sections") else "."))
    if c.get("example"):
        sub(f"Ví dụ — mục {c['example_sec']} của FRD được tách thành {len(c['example'])} yêu cầu")
        for rid, q, typ in c["example"]:
            row(rid, f"{q}   [{typ}]")
        row("", "Mỗi yêu cầu là một điểm test riêng, pass hoặc fail độc lập — ví dụ một yêu cầu có thể lệch với code trong khi yêu cầu kế bên khớp.")
    gap()

    section("3. SHEET TCs — Ý NGHĨA CÁC CỘT")
    row("Dòng 2–7 (đầu trang)", "Tên chức năng + mã Redmine, link FRD, link Redmine, phiên bản FRD đã dùng để đối chiếu.")
    row("Dòng nhóm \"I. …\"", f"Nhóm chức năng (ví dụ \"Xác thực OTP\") và số case trong nhóm. Có {c['n_topics']} nhóm chức năng, xếp theo "
        "thứ tự trong FRD — cùng danh sách với cột Nhóm chức năng ở sheet Ma trận. Case được xếp vào nhóm chứa nhiều yêu cầu FRD của nó "
        "nhất. Nhóm cuối \"Tính năng ngoài FRD\" gồm case kiểm tính năng BA tự thêm. Muốn biết case nào chạy được, xem cột K.")
    row("A · STT trường hợp kiểm thử", "Mã case: TC-* (chức năng), SEC-* (bảo mật), E2E-* (hành trình đầu–cuối). "
        "Dòng ghi \"… (tiếp 2)\" là phần nối tiếp của case ngay trên, do nội dung dài quá giới hạn chiều cao 1 dòng của Excel.")
    row("B · Mục đích kiểm thử", "Tên case + \"(Phủ: …)\" là các mã yêu cầu trong tài liệu BA mà case kiểm (AC-/BR-/VR-/UC-/AF-).")
    row("C · Precondition", "Điều kiện trước khi chạy: dữ liệu seed (SEED-*), cấu hình (SETUP-*), điểm điều khiển (FX-*). "
        "Các mã P-xx trong cột này là điều kiện tiên quyết cũ của bộ QC (viết khi code chưa có) — hiện phần lớn đã có trên staging, "
        "không dùng P-xx để bỏ qua case.")
    row("D · Test data", "Dữ liệu nhập. Case E2E ghi dữ liệu ngay trong từng bước; case có bảng dữ liệu+kết quả chung thì xem cột F.")
    row("E · Các bước thực hiện", "Các bước chạy. Mã EL-xx kèm tên phần tử trong ngoặc, ví dụ EL-41 (Phụ đề màn hình) — "
        "đó là data-testid trên UI (tra ở ui-contract.md).")
    row("F · Kết quả mong muốn", "Mỗi dòng \"• điều cần kiểm → giá trị đúng\". Có thể gồm HTTP status, mã message, trạng thái bản ghi DB, event.")
    row("G · Kết quả hiện tại", "QC chọn: P = Pass, F = Fail, PE = Pending (chưa chạy được / đang chờ). Ô sẽ tự đổi màu.")
    row("H · Mã lỗi", "Ghi mã bug Redmine khi F.")
    row("I · Ghi chú", "Lý do case chưa test được / chờ PO, giả định A-* case dựa vào, và ghi chú gốc của case trong repo.")
    row("J · QC thực hiện", "Tên QC chạy case.")
    row("K · Có test được?", "Phân loại tự động — xem mục 5.")
    row("L · Yêu cầu FRD liên quan", "Các yêu cầu FRD (R-xxx) ở sheet Ma trận mà case này kiểm. \"Không có yêu cầu FRD tương ứng\" = case kiểm "
        "tính năng BA tự thêm, không có trong FRD.")
    row("M · Nguồn (file:dòng)", "Vị trí case gốc trong repo biz-portal-docs (thư mục quality-control/...). Mở khi cần đọc đầy đủ.")
    gap()

    section("4. SHEET MA TRẬN FRD — Ý NGHĨA CÁC CỘT")
    row("R", "Mã của yêu cầu FRD, đánh theo thứ tự xuất hiện trong FRD. Định nghĩa ở mục 2.")
    row("Nhóm chức năng", f"Mảng chức năng của yêu cầu (ví dụ Xác thực OTP, Mã số thuế). Có {c['n_topics']} nhóm — trùng với các mục I, II, III… ở sheet TCs.")
    row("Mã yêu cầu", f"Mã ngắn của một ý cụ thể trong nhóm, dạng <nhóm>.<ý> (ví dụ otp.resend.limit = giới hạn số lần gửi lại OTP). "
        f"Các yêu cầu FRD cùng nói một ý ở nhiều mục có cùng Mã yêu cầu. Có {c['n_keys']} mã. Lọc cột này để gom các yêu cầu trùng ý.")
    row("Mục FRD · Loại", "Vị trí trong FRD (§ mục) và loại yêu cầu (Validation, Flow, Message, Integration…).")
    row("Trích nguyên văn FRD", "Câu gốc trong FRD — đã được script kiểm là có nguyên văn trong FRD, không bị diễn giải lại.")
    row("BA · BA IDs · Chênh lệch BA", "Tài liệu BA (dev sinh) có viết đúng yêu cầu này không, bằng mã nào, lệch ở đâu.")
    row("TA · TA IDs · Chênh lệch TA", "Tài liệu thiết kế kỹ thuật (TA) có thiết kế đúng không, bằng mã nào (API-, ENT-, EL-…), lệch ở đâu.")
    row("Case QC", "Case kiểm yêu cầu này (tối đa 8 mã hiển thị). Bấm vào ô để nhảy tới case đầu tiên ở sheet TCs.")
    row("Case test được ngay", "Những case trong cột Case QC đang ở trạng thái test được (có hoặc không có giả định).")
    row("Tình trạng test", "\"Không có case QC\" · \"Có N case test được\" · \"Chỉ có case chưa test được / chờ PO\".")
    gap()

    section("5. GIÁ TRỊ TRẠNG THÁI")
    sub("Cột K \"Có test được?\" (sheet TCs)")
    row(c["OK"], f"{t[c['OK']]} case — code đã có, case không vướng giả định hay điểm cần PO chốt. Chạy ngay.")
    row(c["WARN"], f"{t[c['WARN']]} case — chạy được, nhưng kết quả đúng dựa trên giả định A-* hoặc câu hỏi chưa chốt (ghi ở cột I). "
        "Fail thì kiểm lại giả định trước khi log bug.")
    row(c["WAIT"], f"{t[c['WAIT']]} case — kiểm hành vi mà BA/TA/code làm KHÁC FRD. Chạy được, nhưng Pass/Fail phụ thuộc PO chọn FRD "
        "hay spec trên Redmine. Không log bug trước khi PO trả lời.")
    row(c["NO"], f"{t[c['NO']]} case — BLOCKED, tính năng chưa có trong code, hoặc chỉ chạy bằng test tự động của dev. "
        "Đánh PE, không tính vào tỉ lệ pass.")
    sub("Cột BA / TA (sheet Ma trận)")
    row("Đủ", "Tài liệu viết đúng nghĩa FRD (giá trị, điều kiện, nhánh, thứ tự).")
    row("Một phần", "Có nhưng thiếu một phần — xem cột Chênh lệch.")
    row("Lệch", "Viết khác FRD (ví dụ FRD resend tối đa 3, tài liệu ghi 5). Cột Chênh lệch trích cả hai bên.")
    row("Thiếu", "Tài liệu không có gì cho yêu cầu này → thường cũng không có case QC.")
    row("N/A", "Không phải yêu cầu kiểm được (nhãn phạm vi, trạng thái tài liệu).")
    gap()

    section("6. CÁCH ĐỌC ĐÚNG — CÂU HỎI THƯỜNG GẶP")
    row("Yêu cầu FRD nào chưa có case?", f"Sheet Ma trận → lọc Tình trạng test = \"Không có case QC\": {c['n_nocase']} yêu cầu, thuộc "
        f"{c['n_nocase_keys']} mã yêu cầu / {c['n_nocase_topics']} nhóm chức năng ({c['top_nocase_topics']}). Đây là danh sách cần bổ sung case "
        "hoặc báo gap — phần lớn là tính năng FRD có nhưng tài liệu BA bỏ.")
    row("Một case kiểm những yêu cầu nào?", "Sheet TCs → cột L \"Yêu cầu FRD liên quan\".")
    row("Một yêu cầu có những case nào?", "Sheet Ma trận → cột Case QC (bấm để nhảy sang TCs). Hoặc ở sheet TCs, dùng Ctrl+F tìm mã R-xxx trong cột L.")
    row(f"Vì sao {c['n_rows']} yêu cầu mà chỉ {c['n_cases']} case?",
        f"Quan hệ nhiều–nhiều: 1 case thường kiểm cùng lúc nhiều yêu cầu (có case phủ {c['max_r_per_case']} yêu cầu), và FRD viết cùng "
        f"một ý ở nhiều mục (mô tả, use case, luồng, yêu cầu chức năng, validation, AC). {c['n_rows_with_case']} yêu cầu có case chỉ "
        f"quy về {c['n_distinct_cases']} case khác nhau. Không cộng số case theo từng yêu cầu.")
    row("Vì sao ma trận không gộp các yêu cầu trùng ý?", "Để mỗi yêu cầu giữ nguyên văn và vị trí trong FRD, đối chiếu ngược được. "
        "Muốn xem theo ý, lọc hoặc sort theo cột Mã yêu cầu; theo mảng chức năng thì lọc cột Nhóm chức năng.")
    row("Case không gắn yêu cầu FRD nào?", f"{c['n_case_no_frd']} case — kiểm tính năng BA tự thêm ngoài FRD. Cột L ghi "
        "\"Không có yêu cầu FRD tương ứng\". Test hay không tuỳ PO chốt phạm vi.")
    row("Dòng \"(tiếp 2)\" là gì?", f"Phần nối tiếp của case ngay trên ({c['n_cont']} dòng như vậy). Kết quả P/F/PE chỉ ghi ở dòng đầu của case.")
    row("Case đã rút?", f"{c['n_withdrawn']} case không đưa vào sheet TCs, xem danh sách ở sheet Tóm tắt.")
    row("Trước khi log bug", "Mở FRD (link ở dòng 3 sheet TCs) và tài liệu gốc (cột M) để xác nhận. Trạng thái Lệch/Thiếu và cột K "
        "do AI đánh giá — đúng phần lớn nhưng không tuyệt đối.")
    gap()

    section("7. BẢNG TRA MÃ")
    for code, desc in [
        ("R-xxx", "Yêu cầu tách từ FRD (sheet Ma trận)."),
        ("TC- / SEC- / E2E- / TF-", "Case chức năng / case bảo mật / hành trình đầu–cuối / luồng kiểm thử."),
        ("AC- / BR- / VR- / UC- / AF-", "Mã trong tài liệu BA: Acceptance Criteria / Business Rule / Validation Rule / Use Case / Alternative Flow."),
        ("API- / ENT- / EVT- / EL- / SC-", "Mã trong tài liệu TA: API / bảng dữ liệu / event / phần tử UI (data-testid) / màn hình."),
        ("CF- / TD- / RSK-", "Mã TA: điểm code đang chạy trái yêu cầu / quyết định kỹ thuật / rủi ro."),
        ("A-xx", "Giả định của bộ case (khi yêu cầu chưa chốt). Case dựa vào giả định được ghi ở cột I."),
        ("P-xx", "Điều kiện tiên quyết cũ của bộ QC — xem ghi chú cột C."),
        ("OQ- / OP- / T- / C- / G-", "Câu hỏi mở: của QC/BA (OQ) · FRD tự ghi TBD (OP/T) · FRD tự mâu thuẫn (C) · FRD thiếu định nghĩa (G)."),
        ("SEED- / SETUP- / FX-", "Dữ liệu seed / cấu hình chuẩn bị / điểm điều khiển (cách đọc OTP, ép trạng thái) trong Precondition."),
    ]:
        row(code, desc)
    if c["ents"]:
        sub("Mã ENT → bảng trong DB (để viết câu SQL kiểm tra)")
        for ent, (entity, table) in sorted(c["ents"].items()):
            row(ent, f"{table}  —  {entity}")
    gap()

    section("8. LƯU Ý KHI DÙNG")
    row("Xuất lại file", "Mỗi lần chạy skill xuất lại sẽ GHI ĐÈ file (bản cũ lưu thành *_backup.xlsx). Nếu đã ghi kết quả test, "
        "đổi tên file hoặc chép kết quả sang trước.")
    row("Dữ liệu DB", "Case assert theo tên bảng/cột trong tài liệu TA (tra mã ENT ở mục 7). File không chứa câu SQL, dữ liệu seed "
        "hay thông tin truy cập DB — xin dev/DevOps.")
    row("Khi PO chốt nguồn chuẩn", "Nếu PO chọn FRD: các case ⏸ kiểm hành vi trái FRD sẽ phải viết lại, và các yêu cầu \"Không có case QC\" "
        "cần bổ sung case. Nếu PO chọn spec Redmine: FRD cần cập nhật, các case ⏸ thành test được.")
    for x in c.get("extra_notes", []):
        row(x.get("title", ""), x.get("text", ""))
    row("Giới hạn", c["limits"])

    for rr in range(1, r[0]):
        a, b = ws.cell(rr, 1).value, ws.cell(rr, 2).value
        lines = max(1, sum(max(1, -(-len(p) // 118)) for p in str(b or "").split("\n")),
                    -(-len(str(a or "")) // 33))
        if ws.row_dimensions[rr].height is None:
            ws.row_dimensions[rr].height = min(409, lines * 15 + 3)
    return ws
